"""Verify the separate workbook against its CSV export and original candidate IDs."""
import csv
import json
from pathlib import Path
import openpyxl

project = Path(__file__).resolve().parents[1]
out = project/'outputs/01a0d7fb-20a1-75c2-83ec-202ecae1a0ed'
wb = openpyxl.load_workbook(out/'review_remaining.xlsx', read_only=False, data_only=False)
with (out/'review.csv').open(encoding='utf-8-sig', newline='') as f:
    rows = list(csv.reader(f))
headers = rows[0]
sheet = wb['Review']
assert [c.value for c in sheet[1]] == headers
assert sheet.max_row == len(rows)
assert len(headers) == 65
label_col = headers.index('Confirmed re-identification')+1
status_col = headers.index('Review status')+1
assert all(sheet.cell(r,label_col).value == 'Unsure' for r in range(2,sheet.max_row+1))
assert all(sheet.cell(r,status_col).value.startswith('pending') for r in range(2,sheet.max_row+1))
assert sheet.freeze_panes and len(sheet.tables) == 1
prior = set(json.loads((project/'data_remaining/inputs/prior_candidate_ids.json').read_text()))
new_ids = set()
with (out/'keyword_hits.csv').open(encoding='utf-8-sig',newline='') as f:
    for row in csv.DictReader(f):
        new_ids.add(row['record_id'])
assert not prior.intersection(new_ids)
assert len(new_ids) == sheet.max_row-1
summary = {row[0].value:row[1].value for row in wb['Summary'].iter_rows(min_row=2)}
assert summary['Time limit'] == 'None'
assert summary['Newly completed files'] == 2
assert summary['New records in completed files'] == 1000000
assert summary['Files still unfinished'] == 2092
assert summary['Unique keyword-matching records'] == len(new_ids)
print(json.dumps({'review_rows':len(new_ids),'sheets':len(wb.sheetnames),
                  'prior_candidate_overlap':0,'newly_scanned_records':1000000,
                  'pending_labels_verified':True,'separate_workbook_verified':True}))
